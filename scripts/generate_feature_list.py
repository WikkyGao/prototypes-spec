#!/usr/bin/env python3
"""从原型解析数据生成功能清单 Excel 文件。

用法：
    python3 scripts/generate_feature_list.py <解析JSON> [输出路径]

输入：parse_axure.py --json 输出的 JSON 文件路径
输出：功能清单.xlsx（主表 + 副表「说明」）
"""
import json
import os
import sys
import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADERS = ['功能编号', '模块', '功能名称', '功能描述', '优先级',
           '所属页面', '关联功能', '备注']
WIDTHS = [12, 18, 22, 40, 8, 18, 18, 30]


def infer_features_for_page(page, start_no):
    """根据页面内容推断功能清单。"""
    features = []
    module = page['pageName']
    notes = page.get('notes', [])
    notes_text = '; '.join(notes) if notes else ''

    # 规则1：页面有文本 → 至少一个"查看"主功能
    if page.get('texts'):
        main = infer_main_feature(page)
        features.append({
            'id': f'F-{start_no:03d}', 'module': module,
            'name': main['name'], 'desc': main['desc'],
            'priority': 'P0', 'page': module, 'related': '',
            'note': notes_text,
        })
        start_no += 1

    # 规则2：从控件类型推断功能
    widget_features = infer_features_from_widgets(page, start_no)
    features.extend(widget_features)
    start_no += len(widget_features)

    # 规则3：从交互事件推断功能
    for evt in page.get('interactions', [])[:5]:
        feat_name = (evt.get('description', '') or '交互')[:20]
        features.append({
            'id': f'F-{start_no:03d}', 'module': module,
            'name': feat_name,
            'desc': f'触发{evt.get("event","")}事件：{evt.get("description","")}',
            'priority': 'P1', 'page': module, 'related': '',
            'note': '',
        })
        start_no += 1

    # 规则4：从备注说明补功能
    for note in notes[:2]:
        if any(kw in note for kw in ['有效期', '验证', '规则', '必须', '只能']):
            features.append({
                'id': f'F-{start_no:03d}', 'module': module,
                'name': f'业务规则：{note[:15]}', 'desc': note,
                'priority': 'P1', 'page': module, 'related': '',
                'note': '来自备注说明',
            })
            start_no += 1

    return features


def infer_main_feature(page):
    """推断页面主功能。"""
    module = page['pageName']
    if '管理' in module:
        return {'name': f'查看{module}列表', 'desc': f'展示{module}的相关信息列表'}
    elif '新增' in module:
        base = module.replace('新增', '')
        return {'name': f'新增{base}', 'desc': f'填写表单新增{base}'}
    elif '详情' in module:
        return {'name': f'查看{module}详情', 'desc': f'展示{module}的详细信息'}
    else:
        return {'name': f'查看{module}', 'desc': f'展示{module}页面内容'}


def infer_features_from_widgets(page, start_no):
    """从控件类型推断功能。"""
    features = []
    module = page['pageName']
    for w in page.get('widgets', [])[:10]:
        w_type = w.get('friendlyType', '') or w.get('type', '')
        label = w.get('label', '') or w.get('name', '')
        if w_type == '表格':
            features.append({
                'id': f'F-{start_no:03d}', 'module': module,
                'name': f'查看{label or "列表"}',
                'desc': f'展示{label or "数据"}列表',
                'priority': 'P0', 'page': module, 'related': '', 'note': '',
            })
            start_no += 1
        elif w_type == '输入框' and label:
            features.append({
                'id': f'F-{start_no:03d}', 'module': module,
                'name': f'填写{label}', 'desc': f'在{label}输入框中填写信息',
                'priority': 'P1', 'page': module, 'related': '', 'note': '',
            })
            start_no += 1
        elif w_type == '树':
            features.append({
                'id': f'F-{start_no:03d}', 'module': module,
                'name': '树形导航', 'desc': '通过树形结构导航筛选内容',
                'priority': 'P1', 'page': module, 'related': '', 'note': '',
            })
            start_no += 1
    return features


def generate_feature_list_xlsx(pages, output_path):
    """生成功能清单 xlsx 文件。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '功能清单'

    # 样式
    header_font = Font(bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='548235', end_color='548235', fill_type='solid')
    module_font = Font(bold=True, size=11, color='375623')
    module_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_align = Alignment(vertical='center', wrap_text=True)

    # 写表头
    for col, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 28

    # 按页面分组写功能
    row = 2
    f_no = 1
    for page in pages:
        module = page['pageName']
        # 模块标题行（合并 A:H）
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row=row, column=1, value=f'【{module}】')
        cell.font = module_font
        cell.fill = module_fill
        cell.alignment = Alignment(horizontal='left', vertical='center')
        row += 1

        # 为该页面推断功能
        features = infer_features_for_page(page, f_no)
        for feat in features:
            for col, key in enumerate(
                ['id', 'module', 'name', 'desc', 'priority',
                 'page', 'related', 'note'], 1
            ):
                cell = ws.cell(row=row, column=col, value=feat.get(key, ''))
                cell.border = border
                cell.alignment = wrap_align
            row += 1
            f_no += 1

    ws.freeze_panes = 'A2'

    # 副表「说明」
    ws2 = wb.create_sheet('说明')
    notes_included = any(p.get('notes') for p in pages)
    meta = [
        ['项目名称', pages[0].get('pageName', '') if pages else ''],
        ['生成日期', datetime.date.today().isoformat()],
        ['数据来源', '原型解析'],
        ['功能总数', f_no - 1],
        ['备注', '备注说明已纳入' if notes_included else '无便签/备注'],
    ]
    for r, (k, v) in enumerate(meta, 1):
        ws2.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws2.cell(row=r, column=2, value=v)
    ws2.column_dimensions['A'].width = 14
    ws2.column_dimensions['B'].width = 40

    wb.save(output_path)


def main():
    if len(sys.argv) < 2:
        print('用法：python3 scripts/generate_feature_list.py <解析JSON> [输出路径]',
              file=sys.stderr)
        sys.exit(1)
    json_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else '功能清单.xlsx'

    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    pages = data.get('pages', [])

    generate_feature_list_xlsx(pages, output_path)
    print(f'✓ 已生成 {output_path}（{len(pages)} 页面）')


if __name__ == '__main__':
    main()
