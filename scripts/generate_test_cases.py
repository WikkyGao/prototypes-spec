#!/usr/bin/env python3
"""从原型解析数据生成测试用例 Excel 文件。

用法：
    python3 scripts/generate_test_cases.py <解析JSON> [输出路径]

输入：parse_axure.py --json 输出的 JSON 文件路径
输出：测试用例.xlsx（主表 + 副表「说明」）
"""
import json
import os
import sys
import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADERS = ['用例编号', '模块', '功能点', '用例标题', '前置条件',
           '测试步骤', '预期结果', '优先级']
WIDTHS = [14, 18, 20, 28, 28, 40, 36, 8]


def infer_test_cases_for_page(page, start_no):
    """根据页面内容推断测试用例。"""
    cases = []
    module = page['pageName']

    # 规则1：有文本内容 → 至少 1 个加载用例
    if page.get('texts'):
        cases.append({
            'id': f'TC-{start_no:03d}', 'module': module,
            'feature': '页面加载', 'title': f'正常加载{module}页',
            'precondition': '已登录系统',
            'steps': f'1. 进入{module}页',
            'expected': '页面正常显示，含原型中的文本内容',
            'priority': 'P0',
        })
        start_no += 1

    # 规则2：有交互事件 → 补交互用例
    for evt in page.get('interactions', [])[:3]:
        cases.append({
            'id': f'TC-{start_no:03d}', 'module': module,
            'feature': '交互',
            'title': f'触发{evt.get("event","")}事件',
            'precondition': f'已进入{module}页',
            'steps': f'1. 触发{evt.get("description","")}相关控件',
            'expected': '按原型定义的交互响应',
            'priority': 'P1',
        })
        start_no += 1

    # 规则3：有备注说明 → 补备注相关用例
    for note in page.get('notes', [])[:2]:
        cases.append({
            'id': f'TC-{start_no:03d}', 'module': module,
            'feature': '业务规则',
            'title': f'验证备注说明：{note[:20]}',
            'precondition': f'已进入{module}页',
            'steps': f'1. 按备注说明操作：{note}',
            'expected': '符合备注说明的业务规则',
            'priority': 'P1',
        })
        start_no += 1

    return cases


def generate_test_cases_xlsx(pages, output_path):
    """生成测试用例 xlsx 文件。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '测试用例'

    # 样式
    header_font = Font(bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    module_font = Font(bold=True, size=11, color='1F4E79')
    module_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_align = Alignment(vertical='center', wrap_text=True)

    # 写表头
    for col, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 28

    # 按页面分组写用例
    row = 2
    tc_no = 1
    for page in pages:
        module = page['pageName']
        # 模块标题行（合并 A:H）
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row=row, column=1, value=f'【{module}】')
        cell.font = module_font
        cell.fill = module_fill
        cell.alignment = Alignment(horizontal='left', vertical='center')
        row += 1

        cases = infer_test_cases_for_page(page, tc_no)
        for case in cases:
            for col, key in enumerate(
                ['id', 'module', 'feature', 'title', 'precondition',
                 'steps', 'expected', 'priority'], 1
            ):
                cell = ws.cell(row=row, column=col, value=case.get(key, ''))
                cell.border = border
                cell.alignment = wrap_align
            row += 1
            tc_no += 1

    ws.freeze_panes = 'A2'

    # 副表「说明」
    ws2 = wb.create_sheet('说明')
    notes_included = any(p.get('notes') for p in pages)
    meta = [
        ['项目名称', pages[0].get('pageName', '') if pages else ''],
        ['生成日期', datetime.date.today().isoformat()],
        ['数据来源', '原型解析'],
        ['用例总数', tc_no - 1],
        ['备注', '备注说明已纳入' if notes_included else '无便签/备注'],
    ]
    for r, (k, v) in enumerate(meta, 1):
        ws2.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws2.cell(row=r, column=2, value=v)
    ws2.column_dimensions['A'].width = 14
    ws2.column_dimensions['B'].width = 40

    wb.save(output_path)


def main():
    if len(sys.argv) < 2:
        print('用法：python3 scripts/generate_test_cases.py <解析JSON> [输出路径]',
              file=sys.stderr)
        sys.exit(1)
    json_path = sys.argv[1]
    output_path = sys.argv[2] if len(sys.argv) > 2 else '测试用例.xlsx'

    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    pages = data.get('pages', [])

    generate_test_cases_xlsx(pages, output_path)
    print(f'✓ 已生成 {output_path}（{len(pages)} 页面）')


if __name__ == '__main__':
    main()
